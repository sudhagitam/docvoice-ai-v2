"""
DocVoice AI – Document Extractor
Handles PDF and DOCX text extraction with robust error handling.
"""

import io
import logging
import re
from pathlib import Path
from typing import Optional

logger = logging.getLogger("docvoice.extractor")


class DocumentExtractor:
    """Extracts plain text from PDF and DOCX files."""

    SUPPORTED = {".pdf", ".docx"}

    # ── Public API ────────────────────────────────────────────────────────────

    def extract(self, file_bytes: bytes, filename: str) -> str:
        """
        Extract text from *file_bytes*.

        Args:
            file_bytes: Raw file content.
            filename:   Original filename (used to detect type).

        Returns:
            Cleaned extracted text string.

        Raises:
            UnsupportedFileTypeError | EmptyDocumentError | CorruptedFileError
        """
        from core.exceptions import (
            CorruptedFileError,
            EmptyDocumentError,
            UnsupportedFileTypeError,
        )

        ext = Path(filename).suffix.lower()
        if ext not in self.SUPPORTED:
            raise UnsupportedFileTypeError(ext)

        logger.info("Extracting text from '%s' (%d bytes)", filename, len(file_bytes))

        try:
            if ext == ".pdf":
                text = self._extract_pdf(file_bytes)
            else:
                text = self._extract_docx(file_bytes)
        except (UnsupportedFileTypeError, EmptyDocumentError, CorruptedFileError):
            raise
        except Exception as exc:
            logger.exception("Unexpected extraction error for '%s'", filename)
            raise CorruptedFileError(str(exc)) from exc

        text = self._clean(text)

        if not text.strip():
            raise EmptyDocumentError()

        logger.info("Extracted %d characters from '%s'", len(text), filename)
        return text

    # ── PDF ──────────────────────────────────────────────────────────────────

    def _extract_pdf(self, data: bytes) -> str:
        """Extract text from a PDF using PyPDF2, with optional pytesseract OCR fallback."""
        from core.exceptions import CorruptedFileError

        try:
            import PyPDF2
        except ImportError as exc:
            raise CorruptedFileError("PyPDF2 not installed") from exc

        try:
            reader = PyPDF2.PdfReader(io.BytesIO(data))
        except Exception as exc:
            raise CorruptedFileError(f"Cannot parse PDF: {exc}") from exc

        if reader.is_encrypted:
            raise CorruptedFileError("Password-protected PDFs are not supported.")

        parts: list[str] = []
        for i, page in enumerate(reader.pages):
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                logger.warning("Failed to extract page %d; skipping.", i + 1)

        text = "\n".join(parts)

        # Fallback to OCR if no text was found
        if not text.strip():
            text = self._ocr_pdf(data)

        return text

    def _ocr_pdf(self, data: bytes) -> str:
        """OCR fallback using pdf2image + pytesseract (optional deps)."""
        try:
            import pytesseract
            from pdf2image import convert_from_bytes

            logger.info("No selectable text found; running OCR …")
            images = convert_from_bytes(data, dpi=200)
            return "\n".join(pytesseract.image_to_string(img) for img in images)
        except ImportError:
            logger.warning("OCR libs not available (pdf2image / pytesseract).")
            return ""
        except Exception as exc:
            logger.warning("OCR failed: %s", exc)
            return ""

    # ── DOCX ─────────────────────────────────────────────────────────────────

    def _extract_docx(self, data: bytes) -> str:
        """Extract text from a DOCX file using python-docx."""
        from core.exceptions import CorruptedFileError

        try:
            from docx import Document
        except ImportError as exc:
            raise CorruptedFileError("python-docx not installed") from exc

        try:
            doc = Document(io.BytesIO(data))
        except Exception as exc:
            raise CorruptedFileError(f"Cannot parse DOCX: {exc}") from exc

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also pull text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)

        return "\n".join(paragraphs)

    # ── Helpers ───────────────────────────────────────────────────────────────

    @staticmethod
    def _clean(text: str) -> str:
        """Normalise whitespace and remove control characters."""
        # Remove non-printable chars except newline/tab
        text = re.sub(r"[^\x09\x0A\x0D\x20-\x7E\u00A0-\uFFFF]", " ", text)
        # Collapse runs of blank lines
        text = re.sub(r"\n{3,}", "\n\n", text)
        # Collapse horizontal whitespace
        text = re.sub(r"[ \t]{2,}", " ", text)
        return text.strip()
